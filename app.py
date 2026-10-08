"""
AI Career Analyzer  (DEMO MODE)
===============================
 
Run with:  streamlit run app.py
 
Question the app answers:
    "Given my CV and the career I want, what should I do next?"
 
--------------------------------------------------------------------------
DEVELOPER SECTION: WHERE REAL AI / BACKENDS CONNECT LATER
--------------------------------------------------------------------------
Target pipeline (NOT implemented yet, the demo uses mock logic):
 
    CV -> Text Extraction -> Structured Profile -> Skill Extraction
       -> Target Career Analysis -> Skill Gap Analysis
       -> Career Roadmap Generation -> Next Action Generation
       -> Streamlit Results
 
File layout (single file, three clearly separated layers):
    PART 1  MOCK AI DATA        - fake knowledge base (roles, skills, effort)
    PART 2  ANALYSIS LOGIC      - the functions the UI calls
    PART 3  UI                  - Streamlit screens; no analysis logic
 
Integration map (function -> future component):
    1. CV parser ................ extract_cv_text()           (basic parser is real)
    2. LLM ...................... generate_career_analysis()   (summary + reasoning)
    3. Skill extraction model ... extract_skills()
    4. Job description analyzer . analyze_skill_gaps()        (job_description arg)
    5. Career knowledge base .... ROLE_TEMPLATES / SKILL_VOCAB (replace with
                                  vector DB, O*NET, or LLM retrieval)
    6. Roadmap generator ........ generate_career_roadmap(), generate_next_actions()
    7. Database ................. none yet; persist `analysis` dicts in
                                  generate_career_analysis() / the UI submit handler
 
Contract: generate_career_analysis() returns a plain dict. As long as a real
implementation returns the same keys, the UI needs no changes.
--------------------------------------------------------------------------
"""
 
from __future__ import annotations
 
import html
import io
import re
from dataclasses import dataclass
 
import streamlit as st
 
DEMO_MODE = True  # TODO: set to False once the real AI pipeline is connected.
 
 
# ==========================================================================
# PART 1: MOCK AI DATA  (everything here is placeholder knowledge)
# ==========================================================================
# TODO: Replace this whole section with a real career knowledge base / LLM.
 
# Skill -> aliases searched for in CV text (lowercase).
SKILL_VOCAB: dict[str, list[str]] = {
    "Python": ["python"],
    "SQL": ["sql", "mysql", "postgresql", "postgres", "sqlite"],
    "Git": ["git", "github", "gitlab"],
    "C++": ["c++"],
    "Java": ["java"],
    "JavaScript": ["javascript", "node.js", "nodejs"],
    "TypeScript": ["typescript"],
    "React": ["react", "react.js"],
    "HTML/CSS": ["html", "css"],
    "Statistics": ["statistics", "statistical", "hypothesis testing", "probability"],
    "Machine Learning": ["machine learning", "scikit-learn", "sklearn", "xgboost", "random forest"],
    "Deep Learning": ["deep learning", "neural network", "neural networks", "pytorch", "tensorflow", "keras"],
    "NLP": ["nlp", "natural language processing", "llm", "llms", "transformers"],
    "Computer Vision": ["computer vision", "opencv", "image processing"],
    "Data Analysis": ["data analysis", "data analytics", "pandas", "numpy", "exploratory data analysis"],
    "Data Visualization": ["data visualization", "tableau", "power bi", "matplotlib", "seaborn", "dashboards"],
    "Data Engineering": ["etl", "data pipeline", "data pipelines", "airflow", "spark"],
    "Excel": ["excel", "spreadsheets"],
    "Docker": ["docker", "containerization"],
    "Kubernetes": ["kubernetes", "k8s"],
    "Cloud": ["aws", "azure", "gcp", "google cloud", "cloud"],
    "CI/CD": ["ci/cd", "github actions", "jenkins", "continuous integration"],
    "MLOps": ["mlops", "mlflow", "model deployment", "model serving"],
    "Linux": ["linux", "bash", "shell scripting"],
    "Networking": ["networking", "tcp/ip", "routing"],
    "Security": ["cybersecurity", "security", "penetration testing", "owasp"],
    "APIs": ["rest api", "restful", "api", "apis", "fastapi", "flask", "django"],
    "System Design": ["system design", "microservices", "distributed systems"],
    "Testing": ["unit testing", "pytest", "test automation", "jest", "testing"],
    "Agile/Scrum": ["agile", "scrum", "kanban"],
    "Project Management": ["project management", "project manager", "stakeholder", "stakeholders"],
    "Communication": ["communication", "presentation", "presentations", "presented", "public speaking", "collaboration", "collaborated"],
    "Leadership": ["leadership", "led a team", "team lead", "mentored", "managed a team"],
    "Problem Solving": ["problem solving", "problem-solving", "analytical"],
    "Product Strategy": ["product strategy", "product management", "product roadmap", "go-to-market"],
    "User Research": ["user research", "user interviews", "usability testing"],
    "UX Design": ["ux", "ui/ux", "user experience", "interaction design"],
    "Figma": ["figma", "sketch", "wireframe", "wireframes", "prototyping"],
    "A/B Testing": ["a/b testing", "ab testing", "experimentation"],
    "Marketing Analytics": ["seo", "google analytics", "campaign", "campaigns", "marketing analytics"],
    "Content Writing": ["copywriting", "content writing", "content marketing"],
}
 
# Role templates: expected skills + importance. `match` = strong title keywords,
# `weak` = low-weight keywords. Any job title can be entered; unknown titles use
# GENERIC_TEMPLATE (plus skills detected in a pasted job description).
ROLE_TEMPLATES: list[dict] = [
    {
        "name": "Machine Learning / AI",
        "match": ["machine learning", "ml", "ai", "artificial intelligence", "data scientist", "deep learning", "nlp", "mlops"],
        "weak": [],
        "skills": {"Python": "Critical", "Machine Learning": "Critical", "Statistics": "High", "Deep Learning": "High",
                   "SQL": "High", "Data Analysis": "High", "Git": "Medium", "Docker": "Medium", "Cloud": "Medium", "MLOps": "Medium"},
    },
    {
        "name": "Data / Business Analyst",
        "match": ["data analyst", "business analyst", "analytics", "business intelligence", "bi", "analyst"],
        "weak": [],
        "skills": {"SQL": "Critical", "Data Analysis": "Critical", "Data Visualization": "Critical", "Excel": "High",
                   "Statistics": "High", "Communication": "High", "Python": "Medium", "A/B Testing": "Medium"},
    },
    {
        "name": "Software / Backend Engineering",
        "match": ["software", "backend", "back-end", "full stack", "fullstack", "programmer"],
        "weak": ["developer", "engineer"],
        "skills": {"Git": "Critical", "APIs": "Critical", "SQL": "High", "Testing": "High", "System Design": "High",
                   "Docker": "Medium", "CI/CD": "Medium", "Cloud": "Medium", "Linux": "Medium"},
    },
    {
        "name": "Frontend / Web Development",
        "match": ["frontend", "front-end", "front end", "web developer", "ui developer", "react"],
        "weak": [],
        "skills": {"JavaScript": "Critical", "HTML/CSS": "Critical", "React": "High", "TypeScript": "High",
                   "Git": "High", "Testing": "Medium", "APIs": "Medium", "UX Design": "Medium"},
    },
    {
        "name": "DevOps / Cloud",
        "match": ["devops", "cloud", "sre", "site reliability", "infrastructure", "platform engineer"],
        "weak": [],
        "skills": {"Linux": "Critical", "Docker": "Critical", "CI/CD": "Critical", "Cloud": "Critical", "Kubernetes": "High",
                   "Git": "High", "Python": "Medium", "Security": "Medium", "Networking": "Medium"},
    },
    {
        "name": "Product Management",
        "match": ["product manager", "product management", "product owner", "pm"],
        "weak": [],
        "skills": {"Product Strategy": "Critical", "Communication": "Critical", "User Research": "High", "Data Analysis": "High",
                   "Agile/Scrum": "High", "Project Management": "High", "A/B Testing": "Medium", "SQL": "Medium"},
    },
    {
        "name": "UX / Product Design",
        "match": ["ux", "ui", "product designer", "designer", "user experience"],
        "weak": [],
        "skills": {"UX Design": "Critical", "Figma": "Critical", "User Research": "Critical", "Communication": "High",
                   "HTML/CSS": "Medium", "Agile/Scrum": "Medium"},
    },
    {
        "name": "Marketing / Growth",
        "match": ["marketing", "seo", "content", "growth", "digital marketing"],
        "weak": [],
        "skills": {"Marketing Analytics": "Critical", "Content Writing": "High", "Data Analysis": "High", "Communication": "High",
                   "A/B Testing": "Medium", "Project Management": "Medium", "Excel": "Medium"},
    },
    {
        "name": "Cybersecurity",
        "match": ["security", "cybersecurity", "penetration", "soc analyst"],
        "weak": [],
        "skills": {"Security": "Critical", "Networking": "Critical", "Linux": "High", "Python": "Medium",
                   "Cloud": "Medium", "Problem Solving": "Medium"},
    },
]
 
GENERIC_TEMPLATE: dict = {
    "name": "General professional role",
    "skills": {"Communication": "High", "Problem Solving": "High", "Project Management": "Medium",
               "Data Analysis": "Medium", "Excel": "Medium", "Leadership": "Medium"},
}
 
IMPORTANCE_RANK = {"Critical": 0, "High": 1, "Medium": 2}
IMPORTANCE_WEIGHT = {"Critical": 3, "High": 2, "Medium": 1}
PRIORITY_LABEL = {"Critical": "Very High", "High": "High", "Medium": "Medium"}
REQUIRED_STATE = {
    "Critical": "Solid practical experience, shown in projects or work",
    "High": "Working proficiency you can demonstrate",
    "Medium": "Basic familiarity",
}
 
# Skill -> (min weeks, max weeks) to reach working level from scratch.
SKILL_EFFORT: dict[str, tuple[int, int]] = {
    "Python": (4, 6), "SQL": (2, 3), "Git": (1, 1), "Statistics": (3, 5), "Machine Learning": (4, 6),
    "Deep Learning": (4, 8), "Docker": (1, 2), "Kubernetes": (3, 4), "Cloud": (3, 4), "CI/CD": (1, 2),
    "MLOps": (2, 4), "Linux": (2, 3), "JavaScript": (4, 6), "TypeScript": (2, 3), "React": (3, 5),
    "HTML/CSS": (2, 3), "Excel": (1, 2), "Data Visualization": (2, 3), "Data Analysis": (3, 4),
    "System Design": (4, 6), "Testing": (1, 2), "APIs": (2, 3), "Figma": (2, 3), "User Research": (2, 4),
}
DEFAULT_EFFORT = (2, 4)
 
# Short "why this matters" notes, used in mock explanations.
SKILL_NOTES: dict[str, str] = {
    "SQL": "Most data-related work starts with getting and shaping data from databases.",
    "Statistics": "It underpins how results are validated and how decisions are justified.",
    "Docker": "Employers expect work to be reproducible and deployable, not just run locally.",
    "Cloud": "Production systems for this kind of role usually run on a cloud platform.",
    "MLOps": "Getting models into production and keeping them healthy is a core part of the job.",
    "Git": "Version control is a baseline expectation for collaborative technical work.",
    "Communication": "Results only matter if you can explain them to non-specialists.",
    "Testing": "Reliable work needs automated checks, and hiring teams look for this habit.",
    "APIs": "Most modern systems are built by connecting services through APIs.",
    "User Research": "Good decisions start from evidence about real user needs.",
}
 
LEVELS = ["Student", "Entry Level", "Junior", "Mid-Level", "Senior"]
 
 
# ==========================================================================
# PART 2: ANALYSIS LOGIC  (UI-independent; swap internals for real AI later)
# ==========================================================================
 
@dataclass
class CVFile:
    """Lightweight container so the CV survives Streamlit reruns."""
    name: str
    data: bytes
 
 
class CVParseError(Exception):
    """Raised with a user-friendly message when a CV can't be read."""
 
 
def extract_cv_text(cv_file: CVFile | None) -> str:
    """Basic real text extraction for PDF/DOCX.
    TODO: integrate a stronger CV parser (layout-aware, OCR for scanned PDFs).
    """
    if cv_file is None:
        return ""
    ext = cv_file.name.lower().rsplit(".", 1)[-1] if "." in cv_file.name else ""
    try:
        if ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(cv_file.data))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if ext == "docx":
            from docx import Document
            doc = Document(io.BytesIO(cv_file.data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells)
            return "\n".join(parts)
    except ImportError as exc:
        raise CVParseError("A CV parsing library is missing. Run `pip install -r requirements.txt`.") from exc
    except Exception as exc:  # corrupted / encrypted file
        raise CVParseError("We couldn't read this file. Please try another PDF or DOCX.") from exc
    raise CVParseError("Unsupported file type. Please upload a PDF or DOCX file.")
 
 
def extract_skills(cv_text: str) -> dict[str, int]:
    """Return {skill: number of mentions}.
    DEMO IMPLEMENTATION: keyword matching.
    TODO: replace with an LLM / skill-extraction model that judges real evidence.
    """
    text = cv_text.lower()
    found: dict[str, int] = {}
    for skill, aliases in SKILL_VOCAB.items():
        count = 0
        for alias in aliases:
            pattern = r"(?<![\w+#])" + re.escape(alias) + r"(?![\w+#])"
            count += len(re.findall(pattern, text))
        if count:
            found[skill] = count
    return found
 
 
def analyze_cv(cv_file: CVFile | None, level: str, background: str) -> dict:
    """Build a structured profile from the CV and the user's notes.
    DEMO IMPLEMENTATION: text extraction is real, the structuring is shallow.
    TODO: integrate real CV parser + AI (education, experience, projects, achievements).
    """
    cv_text = extract_cv_text(cv_file)
    skills = extract_skills(cv_text + "\n" + (background or ""))
    has_cv = bool(cv_text.strip())
    return {
        "has_cv": has_cv,
        "cv_uploaded": cv_file is not None,
        "level": level,
        "background": (background or "").strip(),
        "skills": skills,
        "has_measurable_results": bool(re.search(r"\d+\s?%|\$\s?\d|\b\d+(?:\.\d+)?\s?[xkm]\b", cv_text, re.I)),
    }
 
 
def _match_template(target_role: str) -> dict | None:
    title = target_role.lower()
    best, best_score = None, 0
    for tpl in ROLE_TEMPLATES:
        score = sum(len(kw) for kw in tpl["match"] if re.search(r"\b" + re.escape(kw) + r"\b", title))
        score += sum(1 for kw in tpl["weak"] if re.search(r"\b" + re.escape(kw) + r"\b", title))
        if score > best_score:
            best, best_score = tpl, score
    return best
 
 
def analyze_skill_gaps(profile: dict, target_role: str, job_description: str = "") -> dict:
    """Compare the profile against what the target role expects.
    DEMO IMPLEMENTATION: role templates + skills spotted in the job description.
    TODO: connect a job-description analyzer and a career knowledge base / LLM.
    """
    template = _match_template(target_role)
    expected: dict[str, str] = dict(template["skills"]) if template else {}
    used_jd = False
 
    jd_skills = extract_skills(job_description) if job_description.strip() else {}
    if jd_skills:
        used_jd = True
        for skill, n in jd_skills.items():
            if skill in expected:
                expected[skill] = "Critical"           # named in the real posting
            else:
                expected[skill] = "High" if n >= 2 else "Medium"
    if not expected:
        expected = dict(GENERIC_TEMPLATE["skills"])
 
    strong, improve, missing, gaps = [], [], [], []
    ordered = sorted(expected.items(), key=lambda kv: IMPORTANCE_RANK[kv[1]])
    for skill, importance in ordered:
        mentions = profile["skills"].get(skill, 0)
        if mentions >= 2:
            strong.append(skill)
            continue
        status = "improve" if mentions == 1 else "missing"
        (improve if status == "improve" else missing).append(skill)
        note = SKILL_NOTES.get(skill, f"It is commonly expected for {target_role} roles.")
        if status == "improve":
            current = "Mentioned in your CV, but with little supporting detail"
            explanation = (f"{skill} shows up in your CV, but not with enough context to prove practical experience. "
                           f"{note} Back it up with one concrete project or result before moving on.")
        else:
            current = "No evidence in your CV" if profile["has_cv"] else "Unknown (no CV provided)"
            explanation = (f"{skill} is a {importance.lower()}-importance expectation for this role, and your CV does not "
                           f"show it. {note}")
        gaps.append({
            "skill": skill, "status": status, "importance": importance,
            "current": current, "required": REQUIRED_STATE[importance], "explanation": explanation,
        })
    gaps.sort(key=lambda g: (IMPORTANCE_RANK[g["importance"]], g["status"] != "missing"))
 
    return {
        "expected": expected, "strong": strong, "improve": improve, "missing": missing, "gaps": gaps,
        "role_family": template["name"] if template else GENERIC_TEMPLATE["name"],
        "role_recognised": template is not None, "used_jd": used_jd,
    }
 
 
def _effort_text(skill: str, status: str) -> str:
    lo, hi = SKILL_EFFORT.get(skill, DEFAULT_EFFORT)
    if status == "improve":
        lo, hi = max(1, lo // 2), max(2, (hi + 1) // 2)
    return f"{lo} week" if lo == hi == 1 else f"{lo}–{hi} weeks"
 
 
def generate_career_roadmap(profile: dict, target_role: str, skill_gaps: dict) -> list[dict]:
    """Ordered steps derived from the user's own gaps.
    DEMO IMPLEMENTATION: rule-based.
    TODO: connect LLM roadmap generator (courses, resources, sequencing by dependency).
    """
    gaps = skill_gaps["gaps"]
    core = gaps[:4]
    steps: list[dict] = []
 
    for g in core:
        verb = "Strengthen" if g["status"] == "improve" else "Learn"
        steps.append({
            "title": f"{verb} {g['skill']}",
            "priority": PRIORITY_LABEL[g["importance"]],
            "effort": _effort_text(g["skill"], g["status"]),
            "description": g["explanation"],
            "done_when": f"You can use {g['skill']} on a real task and explain what you did.",
        })
 
    project_skills = [g["skill"] for g in core[:3]] or skill_gaps["strong"][:3]
    skills_txt = ", ".join(project_skills) if project_skills else "the core skills for the role"
    extra = " Publish it with a clear README." if profile["level"] in ("Student", "Entry Level") else ""
    steps.append({
        "title": f"Build one end-to-end project for {target_role}",
        "priority": "Very High", "effort": "2–4 weeks",
        "description": f"Combine {skills_txt} in a single project that mirrors real work for this role. "
                       f"One finished, well-explained project is worth more than several tutorials.{extra}",
        "done_when": "The project is public, documented, and has at least one measurable result.",
    })
 
    if profile["has_cv"] and not profile["has_measurable_results"]:
        cv_desc = "Your CV lists activities but few measurable outcomes. Rewrite key bullets around results, then prepare for common interview questions."
    elif profile["has_cv"]:
        cv_desc = f"Tailor your CV to {target_role} postings, add your new project, and practise explaining your decisions."
    else:
        cv_desc = "Write a one-page CV focused on this role, add your new project, and practise common interview questions."
    steps.append({
        "title": "Update CV, portfolio and interview preparation",
        "priority": "High", "effort": "1–2 weeks",
        "description": cv_desc,
        "done_when": "Your CV matches the role's keywords and you can walk through your project in five minutes.",
    })
    return steps
 
 
def generate_next_actions(profile: dict, skill_gaps: dict) -> list[dict]:
    """Three concrete things to do now.
    DEMO IMPLEMENTATION: rule-based.
    TODO: connect LLM for personalised, resource-specific actions.
    """
    gaps = skill_gaps["gaps"]
    actions: list[dict] = []
    if gaps:
        top = gaps[0]
        verb = "Strengthen" if top["status"] == "improve" else "Start learning"
        actions.append({"title": f"{verb} {top['skill']}",
                        "detail": f"It is your highest-priority gap. Plan roughly {_effort_text(top['skill'], top['status'])}."})
    else:
        actions.append({"title": "Target stretch skills",
                        "detail": "Your CV covers the core expectations. Look at senior postings for the next skills to add."})
 
    focus = ", ".join(g["skill"] for g in gaps[1:3]) or "your strongest skills"
    actions.append({"title": "Build a small project with real data",
                    "detail": f"Use it to demonstrate {focus} and give your CV something concrete to point to."})
 
    if not profile["has_cv"]:
        actions.append({"title": "Create or upload a CV and re-run the analysis",
                        "detail": "Without a CV this analysis is general. A CV makes it specific to you."})
    elif not profile["has_measurable_results"]:
        actions.append({"title": "Add measurable results to your CV",
                        "detail": "For example: what improved, by how much, for how many users."})
    else:
        actions.append({"title": "Tailor your CV to the target role",
                        "detail": "Mirror the language of real job postings and move your most relevant work to the top."})
    return actions
 
 
def _compute_readiness(expected: dict[str, str], gaps_info: dict) -> int:
    total = got = 0.0
    for skill, importance in expected.items():
        w = IMPORTANCE_WEIGHT[importance]
        total += w
        if skill in gaps_info["strong"]:
            got += w
        elif skill in gaps_info["improve"]:
            got += w * 0.5
    return round(100 * got / total) if total else 0
 
 
def _build_summary(role: str, level: str, readiness: int | None, g: dict, profile: dict) -> str:
    """Mock reasoning text. TODO: replace with LLM-written reasoning."""
    if readiness is None:
        return (f"No CV or background details were provided, so this is not a personal assessment. "
                f"Below is what is generally expected for {role}. Add your CV to get a readiness estimate.")
    strong, gaps = g["strong"], g["gaps"]
    if readiness >= 75:
        opening = f"You look well positioned for {role}."
    elif readiness >= 50:
        opening = f"You have a workable base for {role}, with some important gaps to close."
    elif readiness >= 25:
        opening = f"You are at an early stage for {role}, but the path forward is clear."
    else:
        opening = f"Your CV shows limited overlap with {role} so far, which is common at the start of a transition."
    have = (f" Your CV supports {', '.join(strong[:3])}, which align with what the role expects." if strong
            else " Your CV does not yet show clear evidence of the core skills for this role.")
    if gaps:
        names = [x["skill"] for x in gaps[:3]]
        need = f" The most important gaps are {', '.join(names)}. Close these before polishing your applications."
    else:
        need = " No major gaps were found, so focus on projects, measurable results and interview preparation."
    return opening + have + need
 
 
def generate_career_analysis(cv_file: CVFile | None, target_role: str, level: str,
                             background: str = "", job_description: str = "") -> dict:
    """Run the full (demo) pipeline and return everything the UI renders.
 
    TODO: Replace mock career analysis with real LLM pipeline:
      CV -> text extraction -> structured profile -> skill extraction
         -> target career analysis -> skill gap analysis
         -> roadmap generation -> next actions.
    """
    profile = analyze_cv(cv_file, level, background)
    gaps_info = analyze_skill_gaps(profile, target_role, job_description)
    has_evidence = profile["has_cv"] or bool(profile["background"])
    readiness = _compute_readiness(gaps_info["expected"], gaps_info) if has_evidence else None
    return {
        "is_demo": DEMO_MODE,
        "target_role": target_role,
        "level": level,
        "profile": profile,
        "used_jd": gaps_info["used_jd"],
        "role_recognised": gaps_info["role_recognised"],
        "role_family": gaps_info["role_family"],
        "readiness": readiness,
        "summary": _build_summary(target_role, level, readiness, gaps_info, profile),
        "strong": gaps_info["strong"],
        "improve": gaps_info["improve"],
        "missing": gaps_info["missing"],
        "gaps": gaps_info["gaps"],
        "roadmap": generate_career_roadmap(profile, target_role, gaps_info),
        "next_actions": generate_next_actions(profile, gaps_info),
    }
 
 
# ==========================================================================
# PART 3: UI  (renders results only; contains no analysis logic)
# ==========================================================================
 
CSS = """
<style>
.block-container, [data-testid="stMainBlockContainer"] {max-width: 760px; padding-top: 2.5rem;}
#MainMenu, footer {visibility: hidden;}
.hero-title {font-size: 2.1rem; font-weight: 700; text-align: center; letter-spacing: -0.01em; margin-bottom: .2rem;}
.hero-sub {text-align: center; opacity: .7; margin-bottom: 1.6rem; font-size: 1.05rem;}
.pill {display: inline-block; padding: .15rem .6rem; border-radius: 999px; font-size: .78rem;
       border: 1px solid rgba(128,128,128,.45); opacity: .85;}
.chip {display: inline-block; padding: .2rem .7rem; margin: .15rem .25rem .15rem 0; border-radius: 999px;
       font-size: .86rem; border: 1px solid;}
.chip-strong  {background: rgba(34,160,90,.12);  border-color: rgba(34,160,90,.5);}
.chip-improve {background: rgba(224,160,30,.12); border-color: rgba(224,160,30,.55);}
.chip-missing {background: rgba(214,70,70,.10);  border-color: rgba(214,70,70,.5);}
.score {font-size: 3rem; font-weight: 700; line-height: 1;}
.label {font-size: .78rem; text-transform: uppercase; letter-spacing: .06em; opacity: .6;}
.arrow {text-align: center; opacity: .45; font-size: 1.4rem; margin: .1rem 0;}
.muted {opacity: .65; font-size: .85rem;}
div.stButton > button, div[data-testid="stFormSubmitButton"] > button {width: 100%;}
</style>
"""
 
 
def init_state() -> None:
    st.session_state.setdefault("view", "input")      # input | results | roadmap
    st.session_state.setdefault("inputs", {"target_role": "", "level": "Student", "background": "", "job_description": ""})
    st.session_state.setdefault("cv", None)
    st.session_state.setdefault("analysis", None)
 
 
def go(view: str) -> None:
    st.session_state.view = view
 
 
def reset_all() -> None:
    for key in ("inputs", "cv", "analysis", "view"):
        st.session_state.pop(key, None)
 
 
def chips(items: list[str], kind: str) -> str:
    if not items:
        return '<span class="muted">None</span>'
    return "".join(f'<span class="chip chip-{kind}">{html.escape(s)}</span>' for s in items)
 
 
def render_input_view() -> None:
    st.markdown('<div class="hero-title">AI Career Analyzer</div>'
                '<div class="hero-sub">Understand where you stand, what you\'re missing, and what to do next.</div>',
                unsafe_allow_html=True)
    saved = st.session_state.inputs
 
    with st.form("career_form"):
        uploaded = st.file_uploader("Upload your CV", type=["pdf", "docx"], help="PDF or DOCX. Your file is only used for this analysis.")
        if st.session_state.cv and not uploaded:
            st.caption(f"Using previously uploaded CV: {st.session_state.cv.name}. Upload a new file to replace it.")
        target_role = st.text_input("Target Career / Job Title", value=saved["target_role"],
                                    placeholder="e.g. Machine Learning Engineer, Product Manager, UX Designer")
        level = st.selectbox("Current Experience Level", LEVELS, index=LEVELS.index(saved["level"]))
        background = st.text_area("Tell us anything important about your background (optional)", value=saved["background"],
                                  placeholder="Field of study, years of experience, preferred industry, career interests...", height=90)
        job_description = st.text_area("Paste a job description (optional)", value=saved["job_description"],
                                       placeholder="Paste a real job posting for a more accurate analysis.", height=110)
        submitted = st.form_submit_button("Analyze My Career", type="primary")
 
    if submitted:
        role = target_role.strip()
        if not role:
            st.warning("Please enter the career or job title you're aiming for. It's the only required field.")
            return
        cv = st.session_state.cv
        if uploaded is not None:
            if not uploaded.name.lower().endswith((".pdf", ".docx")):
                st.error("Unsupported file type. Please upload a PDF or DOCX file.")
                return
            cv = CVFile(uploaded.name, uploaded.getvalue())
        st.session_state.inputs = {"target_role": role, "level": level, "background": background, "job_description": job_description}
        st.session_state.cv = cv
        try:
            with st.spinner("Analyzing..."):
                st.session_state.analysis = generate_career_analysis(cv, role, level, background, job_description)
        except CVParseError as exc:
            st.error(str(exc))
            return
        except Exception:
            st.error("Something went wrong while analyzing. Please check your inputs and try again.")
            return
        st.session_state.view = "results"
        st.rerun()
 
    if st.session_state.analysis:
        st.button("Back to my last analysis", on_click=go, args=("results",))
 
 
def render_results_view() -> None:
    a = st.session_state.analysis
    if not a:
        go("input")
        st.rerun()
    p = a["profile"]
 
    st.markdown('<span class="pill">Demo Analysis</span>', unsafe_allow_html=True)
    st.title("Your Career Analysis")
    st.markdown(f"**Target role:** {html.escape(a['target_role'])} &nbsp;·&nbsp; **Level:** {a['level']}")
 
    if not p["cv_uploaded"]:
        st.info("No CV was uploaded, so this analysis is general. Upload your CV for a personalised result.")
    elif not p["has_cv"]:
        st.warning("We couldn't find readable text in your CV (it may be a scanned image). Try a text-based PDF or DOCX.")
    if not a["role_recognised"] and not a["used_jd"]:
        st.caption("This title isn't in the demo knowledge base, so general professional expectations were used. "
                   "Pasting a job description improves accuracy.")
    elif a["used_jd"]:
        st.caption("Your pasted job description was used to decide which skills matter most.")
 
    # --- Career fit
    st.header("Career Fit")
    c1, c2 = st.columns([1, 3])
    with c1:
        if a["readiness"] is None:
            st.markdown('<div class="label">Career readiness</div><div class="score">–</div>', unsafe_allow_html=True)
        else:
            st.markdown(f'<div class="label">Career readiness</div><div class="score">{a["readiness"]}%</div>', unsafe_allow_html=True)
    with c2:
        if a["readiness"] is not None:
            st.progress(a["readiness"] / 100)
        st.write(a["summary"])
 
    # --- Skills
    st.header("Skill Analysis")
    st.markdown("**✅ Strong skills** <span class='muted'>supported by your CV</span>", unsafe_allow_html=True)
    st.markdown(chips(a["strong"], "strong"), unsafe_allow_html=True)
    st.markdown("**⚠️ Skills to improve** <span class='muted'>mentioned, but with limited evidence</span>", unsafe_allow_html=True)
    st.markdown(chips(a["improve"], "improve"), unsafe_allow_html=True)
    st.markdown("**❌ Missing skills** <span class='muted'>expected for the role, not found in your CV</span>", unsafe_allow_html=True)
    st.markdown(chips(a["missing"], "missing"), unsafe_allow_html=True)
 
    # --- Gaps
    st.header("What is standing between you and this role?")
    if not a["gaps"]:
        st.success("No major gaps found against the expected skills. Focus on projects, results and interviews.")
    for i, g in enumerate(a["gaps"][:5], start=1):
        with st.container(border=True):
            st.markdown(f"**{i}. {g['skill']}** &nbsp; <span class='pill'>{g['importance']} importance</span>", unsafe_allow_html=True)
            x, y = st.columns(2)
            x.markdown(f"<div class='label'>Current</div>{html.escape(g['current'])}", unsafe_allow_html=True)
            y.markdown(f"<div class='label'>Required</div>{html.escape(g['required'])}", unsafe_allow_html=True)
            st.write(g["explanation"])
 
    # --- Next actions
    st.header("Your Next 3 Actions")
    with st.container(border=True):
        for i, act in enumerate(a["next_actions"], start=1):
            st.markdown(f"**{i}. {act['title']}**")
            st.caption(act["detail"])
 
    st.write("")
    b1, b2 = st.columns(2)
    b1.button("← Edit my inputs", on_click=go, args=("input",))
    b2.button("View My Career Roadmap →", type="primary", on_click=go, args=("roadmap",))
    st.caption("Demo analysis based on simple keyword matching and sample role data. "
               "It is not a real AI assessment yet.")
 
 
def render_roadmap_view() -> None:
    a = st.session_state.analysis
    if not a:
        go("input")
        st.rerun()
 
    st.markdown('<span class="pill">Demo Analysis</span>', unsafe_allow_html=True)
    st.title("Your Career Roadmap")
    st.markdown(f"Steps to become qualified for **{html.escape(a['target_role'])}**, ordered by what will help most.")
 
    steps = a["roadmap"]
    for i, s in enumerate(steps, start=1):
        with st.container(border=True):
            st.markdown(f"<div class='label'>Step {i}</div>", unsafe_allow_html=True)
            st.markdown(f"### {s['title']}")
            st.markdown(f"**Priority:** {s['priority']} &nbsp;·&nbsp; **Estimated effort:** {s['effort']}")
            st.write(s["description"])
            st.caption(f"Done when: {s['done_when']}")
        if i < len(steps):
            st.markdown('<div class="arrow">↓</div>', unsafe_allow_html=True)
 
    st.write("")
    b1, b2 = st.columns(2)
    b1.button("← Back to analysis", on_click=go, args=("results",))
    b2.button("Start a new analysis", on_click=reset_all)
    st.caption("Effort estimates are rough demo values and will vary with your time and starting point.")
 
 
def main() -> None:
    st.set_page_config(page_title="AI Career Analyzer", page_icon="🧭", layout="centered")
    st.markdown(CSS, unsafe_allow_html=True)
    init_state()
    view = st.session_state.view
    if view == "results":
        render_results_view()
    elif view == "roadmap":
        render_roadmap_view()
    else:
        render_input_view()
 
 
main()
 